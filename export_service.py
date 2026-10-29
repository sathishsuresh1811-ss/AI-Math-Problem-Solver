"""
Export Service Module for Excel (XLSX) and CSV reporting.
Generates a multi-tab workbook with professional formatting, frozen headers,
auto-filters, readable column widths, and real Excel calculation formulas.
"""

import os
import csv
from datetime import datetime
from typing import Dict, List, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config import Config
from database import get_all_problems_db, get_all_results_db, get_db
import analytics

# Styling Constants
NAVY_HEADER_FILL = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
TEAL_HEADER_FILL = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
ZEBRA_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
HEADER_FONT = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Segoe UI", size=14, bold=True, color="1E3A8A")
REGULAR_FONT = Font(name="Segoe UI", size=10)
BOLD_FONT = Font(name="Segoe UI", size=10, bold=True)
THIN_BORDER_SIDE = Side(border_style="thin", color="E2E8F0")
BORDER = Border(left=THIN_BORDER_SIDE, right=THIN_BORDER_SIDE, top=THIN_BORDER_SIDE, bottom=THIN_BORDER_SIDE)

def style_header_row(ws, row_idx=1, fill=NAVY_HEADER_FILL):
    """Apply styling to header row."""
    for cell in ws[row_idx]:
        cell.fill = fill
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row_idx].height = 28

def autofit_columns(ws, min_width=12, max_width=50):
    """Auto-fit column widths with padding."""
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if "\n" in val_str:
                val_str = max(val_str.split("\n"), key=len)
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(min_width, min(max_len + 4, max_width))

def generate_excel_workbook(filepath: str = None) -> str:
    """
    Generate comprehensive 10-sheet Excel workbook.
    Returns the absolute path to the generated XLSX file.
    """
    if not filepath:
        os.makedirs(Config.EXPORTS_PATH, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filepath = os.path.join(Config.EXPORTS_PATH, f"AI_Math_Checker_Analysis_{timestamp}.xlsx")

    wb = openpyxl.Workbook()
    
    # Fetch Data
    problems = get_all_problems_db()
    results = get_all_results_db()
    
    # -------------------------------------------------------------
    # Sheet 1: README & Methodology
    # -------------------------------------------------------------
    ws_readme = wb.active
    ws_readme.title = "README & Methodology"
    ws_readme.views.sheetView[0].showGridLines = True
    
    ws_readme["A1"] = "AI Math Problem Explainer and Error Checker — Capstone Project"
    ws_readme["A1"].font = TITLE_FONT
    ws_readme["A2"] = f"Workbook Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    ws_readme["A2"].font = REGULAR_FONT
    
    readme_sections = [
        ("A4", "1. Research Question", "How reliably do free AI tools and prompt conditions solve school/college mathematics problems step by step, and does self-verification improve reasoning accuracy?"),
        ("A7", "2. Experimental Design", "30 Verified Problems (10 Algebra, 10 Probability & Statistics, 10 Calculus) evaluated across 3 Prompt Conditions (Plain, Step-by-Step, Self-Verification)."),
        ("A10", "3. Prompt Conditions", "Plain: Zero-shot direct baseline. Step-by-Step: Chain-of-thought with explicit formula declarations. Self-Verification: Dual-phase derivation followed by independent mathematical checks."),
        ("A13", "4. Accuracy Scoring Rubric", "Final-Answer Accuracy: 1 = Correct, 0 = Incorrect, -1 = Unscorable/Review. Step Accuracy: (Sum of Step Points / Total Step Points) * 100%. SymPy symbolic and numerical equivalence without eval()."),
        ("A16", "5. Workbook Contents", "Sheets: README, Problem Set, Raw Results, Answer Accuracy, Step Accuracy, Error Analysis, Provider Summary, Prompt Comparison, Self-Verification Analysis, Human Review.")
    ]
    for cell_ref, title, body in readme_sections:
        ws_readme[cell_ref] = title
        ws_readme[cell_ref].font = BOLD_FONT
        row_num = int(cell_ref[1:]) + 1
        ws_readme[f"A{row_num}"] = body
        ws_readme[f"A{row_num}"].font = REGULAR_FONT
    ws_readme.column_dimensions["A"].width = 110

    # -------------------------------------------------------------
    # Sheet 2: Problem Set
    # -------------------------------------------------------------
    ws_prob = wb.create_sheet("Problem Set")
    ws_prob.views.sheetView[0].showGridLines = True
    prob_headers = ["Problem ID", "Topic", "Subtopic", "Difficulty", "Problem Statement", "Reference Answer", "Relevant Formula", "Verification Status"]
    ws_prob.append(prob_headers)
    style_header_row(ws_prob, 1)
    
    for row_idx, p in enumerate(problems, start=2):
        ws_prob.append([
            p["id"], p["topic"], p.get("subtopic", ""), p["difficulty"],
            p["statement"], p["final_answer"], p.get("formula", ""), p.get("verification_status", "Verified")
        ])
        if row_idx % 2 == 0:
            for cell in ws_prob[row_idx]:
                cell.fill = ZEBRA_FILL
                cell.font = REGULAR_FONT
                cell.border = BORDER
        else:
            for cell in ws_prob[row_idx]:
                cell.font = REGULAR_FONT
                cell.border = BORDER
    ws_prob.freeze_panes = "A2"
    ws_prob.auto_filter.ref = ws_prob.dimensions
    autofit_columns(ws_prob)

    # -------------------------------------------------------------
    # Sheet 3: Raw Experiment Results
    # -------------------------------------------------------------
    ws_res = wb.create_sheet("Raw Experiment Results")
    ws_res.views.sheetView[0].showGridLines = True
    res_headers = [
        "Result ID", "Problem ID", "Topic", "Difficulty", "Provider", "Model",
        "Prompt Condition", "Parsed Answer", "Reference Answer", "Is Correct",
        "Step Accuracy %", "Latency (ms)", "Is Demo", "Requires Review", "Timestamp"
    ]
    ws_res.append(res_headers)
    style_header_row(ws_res, 1)
    
    for row_idx, r in enumerate(results, start=2):
        ws_res.append([
            r["id"], r["problem_id"], r["topic"], r["difficulty"], r["provider"], r["model"],
            r["prompt_condition"], r.get("parsed_answer", ""), r.get("reference_answer", ""),
            r.get("is_correct", 0), r.get("step_accuracy", 0.0) / 100.0, r.get("latency_ms", 0),
            "Demo" if r.get("is_demo") == 1 else "Live", "Yes" if r.get("requires_human_review") == 1 else "No",
            r.get("timestamp", "")
        ])
        fill_to_use = ZEBRA_FILL if row_idx % 2 == 0 else PatternFill(fill_type=None)
        for cell in ws_res[row_idx]:
            if fill_to_use.fill_type:
                cell.fill = fill_to_use
            cell.font = REGULAR_FONT
            cell.border = BORDER
        # Percentage format on column K (Step Accuracy)
        ws_res[f"K{row_idx}"].number_format = "0.0%"
        
    ws_res.freeze_panes = "A2"
    ws_res.auto_filter.ref = ws_res.dimensions
    autofit_columns(ws_res)

    # -------------------------------------------------------------
    # Sheet 4: Answer Accuracy
    # -------------------------------------------------------------
    ws_acc = wb.create_sheet("Answer Accuracy")
    ws_acc.views.sheetView[0].showGridLines = True
    ws_acc["A1"] = "Overall Accuracy Summary"
    ws_acc["A1"].font = TITLE_FONT
    
    acc_headers = ["Metric", "Excel Formula / Value", "Notes"]
    ws_acc.append([])
    ws_acc.append(acc_headers)
    style_header_row(ws_acc, 3, TEAL_HEADER_FILL)
    
    total_rows = len(results)
    if total_rows > 0:
        last_row = total_rows + 1
        summary_rows = [
            ("Total Evaluated Runs", f"=COUNTA('Raw Experiment Results'!A2:A{last_row})", "Total completed experimental trials"),
            ("Correct Final Answers", f"=COUNTIF('Raw Experiment Results'!J2:J{last_row}, 1)", "Runs where final answer was verified correct"),
            ("Incorrect Final Answers", f"=COUNTIF('Raw Experiment Results'!J2:J{last_row}, 0)", "Runs with mathematical errors in answer"),
            ("Unscorable / Review", f"=COUNTIF('Raw Experiment Results'!J2:J{last_row}, -1)", "Ambiguous or unparseable runs"),
            ("Final-Answer Accuracy", f"=COUNTIF('Raw Experiment Results'!J2:J{last_row}, 1) / (COUNTIF('Raw Experiment Results'!J2:J{last_row}, 1) + COUNTIF('Raw Experiment Results'!J2:J{last_row}, 0))", "Normalized percentage of scorable runs"),
            ("Mean Step Accuracy", f"=AVERAGE('Raw Experiment Results'!K2:K{last_row})", "Mean normalized intermediate reasoning score"),
            ("Average Latency (ms)", f"=AVERAGE('Raw Experiment Results'!L2:L{last_row})", "Mean response latency across all backends")
        ]
        for row_idx, (metric, formula, notes) in enumerate(summary_rows, start=4):
            ws_acc.append([metric, formula, notes])
            ws_acc[f"A{row_idx}"].font = BOLD_FONT
            ws_acc[f"B{row_idx}"].font = REGULAR_FONT
            ws_acc[f"C{row_idx}"].font = REGULAR_FONT
            for cell in ws_acc[row_idx]:
                cell.border = BORDER
            if metric in ["Final-Answer Accuracy", "Mean Step Accuracy"]:
                ws_acc[f"B{row_idx}"].number_format = "0.0%"
            elif metric == "Average Latency (ms)":
                ws_acc[f"B{row_idx}"].number_format = "0"
    autofit_columns(ws_acc)

    # -------------------------------------------------------------
    # Sheet 5: Step Accuracy
    # -------------------------------------------------------------
    ws_step = wb.create_sheet("Step Accuracy")
    ws_step.views.sheetView[0].showGridLines = True
    ws_step.append(["Result ID", "Problem ID", "Step Number", "Step Description", "Status", "Score", "Feedback"])
    style_header_row(ws_step, 1)
    
    conn = get_db()
    step_rows = conn.execute("""
    SELECT se.*, r.problem_id
    FROM step_evaluations se
    JOIN experiment_results r ON se.result_id = r.id
    ORDER BY r.timestamp DESC, se.step_number ASC
    LIMIT 1000
    """).fetchall()
    conn.close()
    
    for row_idx, s in enumerate(step_rows, start=2):
        ws_step.append([
            s["result_id"], s["problem_id"], s["step_number"],
            s["step_description"], s["status"], s["score"], s["feedback"]
        ])
        for cell in ws_step[row_idx]:
            cell.font = REGULAR_FONT
            cell.border = BORDER
    ws_step.freeze_panes = "A2"
    ws_step.auto_filter.ref = ws_step.dimensions
    autofit_columns(ws_step)

    # -------------------------------------------------------------
    # Sheet 6: Error Analysis
    # -------------------------------------------------------------
    ws_err = wb.create_sheet("Error Analysis")
    ws_err.views.sheetView[0].showGridLines = True
    ws_err.append(["Error ID", "Result ID", "Problem ID", "Category", "Severity", "Evidence"])
    style_header_row(ws_err, 1)
    
    conn = get_db()
    err_rows = conn.execute("""
    SELECT ec.*, r.problem_id
    FROM error_classifications ec
    JOIN experiment_results r ON ec.result_id = r.id
    ORDER BY ec.id ASC
    """).fetchall()
    conn.close()
    
    for row_idx, e in enumerate(err_rows, start=2):
        ws_err.append([
            e["id"], e["result_id"], e["problem_id"],
            e["category"], e["severity"], e["evidence"]
        ])
        for cell in ws_err[row_idx]:
            cell.font = REGULAR_FONT
            cell.border = BORDER
    ws_err.freeze_panes = "A2"
    ws_err.auto_filter.ref = ws_err.dimensions
    autofit_columns(ws_err)

    # -------------------------------------------------------------
    # Sheet 7: Provider Summary
    # -------------------------------------------------------------
    ws_prov = wb.create_sheet("Provider Summary")
    ws_prov.views.sheetView[0].showGridLines = True
    prov_summary = analytics.compute_by_dimension(results, "provider")
    ws_prov.append(["Provider Name", "Total Runs", "Live Runs", "Demo Runs", "Correct Answers", "Final-Answer Accuracy", "Mean Step Accuracy", "Mean Latency (ms)"])
    style_header_row(ws_prov, 1)
    
    for row_idx, p in enumerate(prov_summary, start=2):
        ws_prov.append([
            p["label"], p["total_runs"], p["live_runs"], p["demo_runs"], p["correct_count"],
            p["final_accuracy"] / 100.0, p["mean_step_accuracy"] / 100.0, p["mean_latency_ms"]
        ])
        for cell in ws_prov[row_idx]:
            cell.font = REGULAR_FONT
            cell.border = BORDER
        ws_prov[f"F{row_idx}"].number_format = "0.0%"
        ws_prov[f"G{row_idx}"].number_format = "0.0%"
    autofit_columns(ws_prov)

    # -------------------------------------------------------------
    # Sheet 8: Prompt Comparison
    # -------------------------------------------------------------
    ws_prompt = wb.create_sheet("Prompt Comparison")
    ws_prompt.views.sheetView[0].showGridLines = True
    prompt_comp = analytics.compute_prompt_comparison(results)
    ws_prompt.append(["Prompt Condition", "Total Trials", "Correct Answers", "Final-Answer Accuracy", "Mean Step Accuracy", "Accuracy Delta vs Plain (pp)"])
    style_header_row(ws_prompt, 1)
    
    conditions_data = [
        ("Plain Prompt", prompt_comp["plain"], 0.0),
        ("Step-by-Step Prompt", prompt_comp["step_by_step"], prompt_comp["deltas"]["step_vs_plain_accuracy_pp"]),
        ("Self-Verification Prompt", prompt_comp["self_verification"], prompt_comp["deltas"]["verif_vs_plain_accuracy_pp"])
    ]
    for row_idx, (lbl, data, delta) in enumerate(conditions_data, start=2):
        ws_prompt.append([
            lbl, data["total_runs"], data["correct_count"],
            data["final_accuracy"] / 100.0, data["mean_step_accuracy"] / 100.0,
            f"{delta:+.1f} pp"
        ])
        for cell in ws_prompt[row_idx]:
            cell.font = REGULAR_FONT
            cell.border = BORDER
        ws_prompt[f"D{row_idx}"].number_format = "0.0%"
        ws_prompt[f"E{row_idx}"].number_format = "0.0%"
    autofit_columns(ws_prompt)

    # -------------------------------------------------------------
    # Sheet 9: Self-Verification Analysis
    # -------------------------------------------------------------
    ws_verif = wb.create_sheet("Self-Verification Analysis")
    ws_verif.views.sheetView[0].showGridLines = True
    verif_data = analytics.compute_self_verification_analysis(results)
    ws_verif["A1"] = "Self-Verification Before / After Comparative Analysis"
    ws_verif["A1"].font = TITLE_FONT
    
    ws_verif.append([])
    ws_verif.append(["Verification Metric", "Observed Value", "Analytical Interpretation"])
    style_header_row(ws_verif, 3, TEAL_HEADER_FILL)
    
    v_metrics = [
        ("Total Verification Trials", verif_data["total_verification_runs"], "Self-verification condition test runs"),
        ("Pre-Verification Accuracy", f"{verif_data['pre_verification_accuracy']:.1f}%", "Accuracy prior to independent check"),
        ("Post-Verification Accuracy", f"{verif_data['post_verification_accuracy']:.1f}%", "Accuracy after independent check and correction"),
        ("Net Percentage Point Delta", f"{verif_data['percentage_point_change']:+.1f} pp", "Direct percentage point shift (+/-)"),
        ("Relative Improvement", f"{verif_data['relative_pct_improvement']:+.1f}%", "Relative percentage change"),
        ("Errors Fixed by Verification", verif_data["fixed_count"], "Cases where initial incorrect answer was caught and corrected"),
        ("Errors Introduced by Verification", verif_data["introduced_count"], "Cases where verification broke a correct initial answer"),
        ("Unchanged Accuracy", verif_data["unchanged_count"], "Cases where verification confirmed initial status")
    ]
    for row_idx, (m, v, interp) in enumerate(v_metrics, start=4):
        ws_verif.append([m, v, interp])
        ws_verif[f"A{row_idx}"].font = BOLD_FONT
        ws_verif[f"B{row_idx}"].font = REGULAR_FONT
        ws_verif[f"C{row_idx}"].font = REGULAR_FONT
        for cell in ws_verif[row_idx]:
            cell.border = BORDER
    autofit_columns(ws_verif)

    # -------------------------------------------------------------
    # Sheet 10: Human Review
    # -------------------------------------------------------------
    ws_rev = wb.create_sheet("Human Review")
    ws_rev.views.sheetView[0].showGridLines = True
    ws_rev.append(["Review ID", "Result ID", "Reviewer Name", "Reviewed At", "Original Score", "Corrected Score", "Notes"])
    style_header_row(ws_rev, 1)
    
    conn = get_db()
    rev_rows = conn.execute("SELECT * FROM human_reviews ORDER BY reviewed_at DESC").fetchall()
    conn.close()
    for row_idx, r in enumerate(rev_rows, start=2):
        ws_rev.append([
            r["id"], r["result_id"], r["reviewer_name"], r["reviewed_at"],
            r["original_is_correct"], r["corrected_is_correct"], r["notes"]
        ])
        for cell in ws_rev[row_idx]:
            cell.font = REGULAR_FONT
            cell.border = BORDER
    autofit_columns(ws_rev)

    wb.save(filepath)
    return filepath

def export_results_csv(filepath: str = None) -> str:
    """Export raw experiment results to CSV."""
    if not filepath:
        os.makedirs(Config.EXPORTS_PATH, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filepath = os.path.join(Config.EXPORTS_PATH, f"experiment_results_{timestamp}.csv")

    results = get_all_results_db()
    if not results:
        fieldnames = ["id", "problem_id", "provider", "model", "prompt_condition", "is_correct", "step_accuracy"]
    else:
        fieldnames = list(results[0].keys())

    with open(filepath, mode="w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for r in results:
            writer.writerow(r)

    return filepath

def export_problems_csv(filepath: str = None) -> str:
    """Export problem set to CSV."""
    if not filepath:
        os.makedirs(Config.EXPORTS_PATH, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filepath = os.path.join(Config.EXPORTS_PATH, f"problem_set_{timestamp}.csv")

    problems = get_all_problems_db()
    fieldnames = ["id", "topic", "subtopic", "difficulty", "statement", "final_answer", "formula", "verification_status"]

    with open(filepath, mode="w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        for p in problems:
            writer.writerow(p)

    return filepath
